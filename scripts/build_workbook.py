#!/usr/bin/env python3
"""Merge the per-file inventory CSVs into one browsable Excel workbook.

    python scripts/build_workbook.py

Reads every catalog/file_inventory_*.csv and writes catalog/seismic_datasets.xlsx
with a "Start Here" tab that answers the only question most people actually have:
which datasets can I train on today.

Requires: pandas, openpyxl
"""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
CATALOG_DIR = ROOT / "catalog"
OUT = CATALOG_DIR / "seismic_datasets.xlsx"

# Header styling
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=10)
# ml_ready colour scale: 5 = ready now, 1 = catalogued for completeness
READY_FILL = {
    5: PatternFill("solid", fgColor="C6EFCE"),
    4: PatternFill("solid", fgColor="E2EFDA"),
    3: PatternFill("solid", fgColor="FFF2CC"),
    2: PatternFill("solid", fgColor="FCE4D6"),
    1: PatternFill("solid", fgColor="F2F2F2"),
}
NUMERIC = {"size_bytes", "n_files", "ml_ready", "sample_rate_ms"}


def human(n) -> str:
    try:
        n = float(n)
    except (TypeError, ValueError):
        return ""
    if n <= 0:
        return ""
    for unit in ("B", "KB", "MB", "GB", "TB", "PB"):
        if abs(n) < 1000:
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.2f} {unit}"
        n /= 1000
    return f"{n:.2f} EB"


def load() -> pd.DataFrame:
    csvs = sorted(CATALOG_DIR.glob("file_inventory_*.csv"))
    if not csvs:
        sys.exit(f"no file_inventory_*.csv found in {CATALOG_DIR}")
    frames = []
    for c in csvs:
        try:
            df = pd.read_csv(c, dtype=str, keep_default_na=False)
        except Exception as e:  # a half-written file from a killed agent
            print(f"  !! skipping {c.name}: {e}", file=sys.stderr)
            continue
        df["_source"] = c.stem.replace("file_inventory_", "")
        frames.append(df)
        print(f"  loaded {c.name}: {len(df)} rows")
    if not frames:
        sys.exit("every CSV failed to parse")
    df = pd.concat(frames, ignore_index=True)

    for col in NUMERIC:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    # Recompute size_human from bytes so the two can never disagree
    if "size_bytes" in df.columns:
        df["size_human"] = df["size_bytes"].map(human)
    return df


def autosize(ws, df: pd.DataFrame, cap: int = 60) -> None:
    for i, col in enumerate(df.columns, start=1):
        longest = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)])
        ws.column_dimensions[get_column_letter(i)].width = min(longest + 2, cap)


def write_sheet(writer, name: str, df: pd.DataFrame, note: str = "") -> None:
    if df.empty:
        return
    df.to_excel(writer, sheet_name=name[:31], index=False)
    ws = writer.sheets[name[:31]]

    for cell in ws[1]:
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autosize(ws, df)

    # Shade rows by ML readiness so the eye lands on the 5s
    if "ml_ready" in df.columns:
        col = list(df.columns).index("ml_ready") + 1
        for r, val in enumerate(df["ml_ready"], start=2):
            try:
                fill = READY_FILL.get(int(val))
            except (TypeError, ValueError):
                continue
            if fill:
                ws.cell(row=r, column=col).fill = fill

    # Right-align and thousands-separate the byte column
    if "size_bytes" in df.columns:
        col = list(df.columns).index("size_bytes") + 1
        for r in range(2, len(df) + 2):
            c = ws.cell(row=r, column=col)
            c.number_format = "#,##0"
            c.alignment = Alignment(horizontal="right")

    if note:
        ws.insert_rows(1)
        ws["A1"] = note
        ws["A1"].font = Font(italic=True, size=9, color="555555")
        ws.freeze_panes = "A3"


def main() -> int:
    print(f"reading {CATALOG_DIR}...")
    df = load()
    print(f"\ntotal rows: {len(df)}")

    total = df["size_bytes"].sum(skipna=True) if "size_bytes" in df else 0
    measured = int(df["size_bytes"].notna().sum()) if "size_bytes" in df else 0
    print(f"measured rows: {measured}   total catalogued: {human(total)}")

    sort_cols = [c for c in ("ml_ready", "size_bytes") if c in df.columns]
    ranked = df.sort_values(sort_cols, ascending=[False] * len(sort_cols)) \
        if sort_cols else df

    with pd.ExcelWriter(OUT, engine="openpyxl") as writer:
        # Start Here: only what you can act on immediately
        if "ml_ready" in df.columns:
            start = ranked[ranked["ml_ready"] >= 4]
            write_sheet(
                writer, "Start Here", start,
                "Datasets scoring 4-5 on ML-readiness: permissive licence, direct "
                "download. Score 5 = labels included. Sorted by readiness, then size.",
            )

        write_sheet(writer, "All Files", ranked,
                    "Every catalogued artefact. size_bytes blank = not measured; "
                    "never estimated. Filter on 'product' to isolate angle stacks.")

        # Per-dataset rollup
        if {"dataset_id", "size_bytes"} <= set(df.columns):
            agg = {"size_bytes": "sum", "file_name": "count"}
            rollup = df.groupby(
                [c for c in ("dataset_id", "dataset_name", "country", "basin")
                 if c in df.columns], dropna=False
            ).agg(agg).reset_index()
            rollup = rollup.rename(columns={"file_name": "n_artefacts",
                                            "size_bytes": "total_bytes"})
            rollup["total_human"] = rollup["total_bytes"].map(human)
            rollup = rollup.sort_values("total_bytes", ascending=False)
            write_sheet(writer, "By Dataset", rollup,
                        "One row per dataset. total_bytes sums only measured files, "
                        "so it is a floor, not a total.")

        # Post-stack + angle stacks — the ML-relevant image-domain subset
        if "domain" in df.columns:
            ps = ranked[ranked["domain"].str.contains("post-stack", case=False,
                                                      na=False)]
            write_sheet(writer, "Post-stack", ps,
                        "Image-domain volumes: the substrate for interpretation ML "
                        "(fault, facies, salt, horizon). Includes angle stacks.")

        if "has_labels" in df.columns:
            lab = ranked[ranked["has_labels"].str.lower().isin(["yes", "partial"])]
            write_sheet(writer, "Labeled", lab,
                        "Artefacts shipping labels. Check label_type and "
                        "label_format before committing to a task.")

        # Licence risk: NonCommercial and unclear terms in one place
        if "commercial_ok" in df.columns:
            risk = ranked[~ranked["commercial_ok"].str.lower().eq("yes")]
            write_sheet(writer, "Licence Watch", risk,
                        "commercial_ok is not 'yes'. OpenFWI data is CC BY-NC-SA "
                        "while its code is BSD-3; FaultSeg3D is research-use-only.")

    print(f"\nwrote {OUT}")
    print("tabs: Start Here | All Files | By Dataset | Post-stack | Labeled | "
          "Licence Watch")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
